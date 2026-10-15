import openpyxl
from openpyxl.styles import Font, PatternFill

# create workbook and select active sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Semester_Results"

# add headers
headers = [
    "Roll No",
    "Student Name",
    "Lab Marks (out of 50)",
    "Theory Marks (out of 50)",
    "Total Marks (out of 100)",
    "Result"
]
ws.append(headers)

# take number of students input
n = int(input("Enter number of students: "))

# taking student details one by one
for i in range(n):
    print(f"\nEnter details for Student {i + 1}:")
    roll = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    lab = float(input("Enter Lab Marks (out of 50): "))
    theory = float(input("Enter Theory Marks (out of 50): "))

    total = lab + theory
    if total >= 40:
        result = "Pass"
    else:
        result = "Fail"

    # write to excel row
    ws.append([roll, name, lab, theory, total, result])

# styling header row
bold_font = Font(bold=True)
gray_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = bold_font
    cell.fill = gray_fill

# calculate class average
total_sum = 0
for row in range(2, 2 + n):
    total_sum += ws.cell(row=row, column=5).value

average_score = total_sum / n

# add summary row at the bottom
summary_row = 2 + n
ws.cell(row=summary_row, column=2, value="Class Average").font = bold_font
ws.cell(row=summary_row, column=5, value=average_score).font = bold_font

# save workbook
wb.save("AIML_Batch_Grades.xlsx")
print("\nGradebook generated and saved as AIML_Batch_Grades.xlsx")
